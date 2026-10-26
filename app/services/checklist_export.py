# -*- coding: utf-8 -*-
"""Proje kontrol listesi dışa aktarma: Excel (ProjeKontrol.xlsx düzeni) ve PDF.

PDF için Türkçe karakterler (ğ ş İ ı ö ü ç) gömülü DejaVu Sans ile basılır; reportlab'ın
standart Helvetica fontu bu harfleri göstermez.
"""
from __future__ import annotations

import os
from io import BytesIO

from app.models.project_checklist import STATUS_APPROVED, STATUS_DONE

_FONT_DIR = os.path.join(os.path.dirname(__file__), "..", "utils", "fonts")

STATUS_LABEL = {
    STATUS_APPROVED: "Onaylandı",
    STATUS_DONE: "Onay bekliyor",
}


def _status_text(item) -> str:
    return STATUS_LABEL.get(item.status, "Bekliyor")


def _person(u) -> str:
    if not u:
        return ""
    return f"{(u.first_name or '').strip()} {(u.last_name or '').strip()}".strip() or (u.email or "")


def _responsible(item) -> str:
    """Atamalar: kişi adları, süreç adı, 'Proje ekibi' (virgülle)."""
    from app.services.project_checklist_service import assignment_labels
    return ", ".join(a["label"] for a in assignment_labels(item.project, item))


def _fmt(dt) -> str:
    return dt.strftime("%d.%m.%Y %H:%M") if dt else ""


def build_xlsx(project, items: list, progress: dict) -> bytes:
    """Orijinal ProjeKontrol düzeni: SÜREÇ | grup | ADIM | SORUMLU | DURUM (+ yapan/onaylayan sütunları)."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kontrol Listesi"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(wrap_text=True, vertical="center", horizontal="center")

    ws.merge_cells("A1:I1")
    ws["A1"] = f"{project.name} — KONTROL LİSTESİ"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:I2")
    done = " · TAMAMLANDI " + _fmt(project.completed_at) if project.completed_at else ""
    ws["A2"] = f"İlerleme: %{progress['pct']} ({progress['approved']}/{progress['total']} madde onaylı){done}"
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["SÜREÇ", "", "ADIM", "SORUMLU", "DURUM", "Yapan", "Yapma zamanı", "Onaylayan", "Onay zamanı"]
    hr = 3
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(hr, c, h)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.border = border
        cell.alignment = center
    ws.merge_cells(start_row=hr, start_column=1, end_row=hr, end_column=2)

    r = hr + 1
    first_row = r
    for it in items:
        ws.cell(r, 1, it.stage)
        ws.cell(r, 2, it.item_group or "")
        ws.cell(r, 3, it.text)
        ws.cell(r, 4, _responsible(it))
        mark = "☑" if it.status == STATUS_APPROVED else "☐"
        ws.cell(r, 5, f"{mark} {_status_text(it)}")
        ws.cell(r, 6, _person(it.done_user))
        ws.cell(r, 7, _fmt(it.done_at))
        ws.cell(r, 8, _person(it.approved_user))
        ws.cell(r, 9, _fmt(it.approved_at))
        for c in range(1, 10):
            ws.cell(r, c).border = border
            ws.cell(r, c).alignment = wrap
        r += 1

    # Aşama / alt grup hücrelerini ardışık satırlar boyunca birleştir (orijinal dosyadaki gibi)
    def runs(key):
        out, start = [], 0
        for i in range(1, len(items) + 1):
            if i == len(items) or key(items[i]) != key(items[start]):
                out.append((start, i - 1))
                start = i
        return out

    for s0, e0 in (runs(lambda it: it.stage) if items else []):
        r0, r1 = first_row + s0, first_row + e0
        if all(not it.item_group for it in items[s0:e0 + 1]):
            if r1 > r0:
                ws.merge_cells(start_row=r0, start_column=1, end_row=r1, end_column=2)
            else:
                ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=2)
        elif r1 > r0:
            ws.merge_cells(start_row=r0, start_column=1, end_row=r1, end_column=1)
        for gs, ge in runs(lambda it: (it.stage, it.item_group)):
            if gs < s0 or ge > e0 or not items[gs].item_group or ge == gs:
                continue
            ws.merge_cells(start_row=first_row + gs, start_column=2, end_row=first_row + ge, end_column=2)
    for rr in range(hr + 1, r):  # birleşik hücre kenarlıkları
        for c in (1, 2):
            ws.cell(rr, c).border = border

    for col, w in zip(range(1, 10), (22, 14, 70, 26, 18, 18, 17, 18, 17)):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = ws.cell(hr + 1, 1)
    ws.sheet_view.showGridLines = False

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(project, items: list, progress: dict, tenant_name: str | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", os.path.join(_FONT_DIR, "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", os.path.join(_FONT_DIR, "DejaVuSans-Bold.ttf")))
        pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold")

    body = ParagraphStyle("ck_body", fontName="DejaVu", fontSize=8, leading=10.5)
    small = ParagraphStyle("ck_small", parent=body, fontSize=7, textColor=colors.HexColor("#475569"))
    head = ParagraphStyle("ck_head", parent=body, fontName="DejaVu-Bold", fontSize=8, textColor=colors.white)
    title = ParagraphStyle("ck_title", parent=body, fontName="DejaVu-Bold", fontSize=14, leading=18)
    meta = ParagraphStyle("ck_meta", parent=body, fontSize=9, textColor=colors.HexColor("#475569"))

    def P(text, style=body):
        return Paragraph(escape(str(text or "")).replace("\n", "<br/>"), style)

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm, title=f"{project.name} — Kontrol Listesi")
    flow = []
    if tenant_name:
        flow.append(P(tenant_name, meta))
    flow.append(P(f"{project.name} — Kontrol Listesi", title))
    done = f" · Tamamlandı {_fmt(project.completed_at)}" if project.completed_at else ""
    flow.append(P(f"İlerleme: %{progress['pct']} ({progress['approved']}/{progress['total']} madde onaylı){done}", meta))
    flow.append(Spacer(1, 0.4 * cm))

    rows = [[P("Aşama / Grup", head), P("Adım", head), P("Sorumlu", head), P("Durum", head), P("Yapan / Onaylayan", head)]]
    styles = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a4d80")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    last_stage = None
    for it in items:
        if it.stage != last_stage:
            rows.append([P(it.stage, ParagraphStyle("st", parent=body, fontName="DejaVu-Bold")), "", "", "", ""])
            ri = len(rows) - 1
            styles += [("SPAN", (0, ri), (-1, ri)), ("BACKGROUND", (0, ri), (-1, ri), colors.HexColor("#DDEBF7"))]
            last_stage = it.stage
        mark = "☑" if it.status == STATUS_APPROVED else "☐"
        who = "; ".join(x for x in (
            f"Yapan: {_person(it.done_user)} ({_fmt(it.done_at)})" if it.done_user else "",
            f"Onay: {_person(it.approved_user)} ({_fmt(it.approved_at)})" if it.approved_user else "",
        ) if x)
        rows.append([P(it.item_group or "", small), P(it.text), P(_responsible(it)),
                     P(f"{mark} {_status_text(it)}"), P(who, small)])
    table = Table(rows, colWidths=[3 * cm, 10.2 * cm, 4.2 * cm, 3 * cm, 6.2 * cm], repeatRows=1)
    table.setStyle(TableStyle(styles))
    flow.append(table)
    doc.build(flow)
    return buf.getvalue()
