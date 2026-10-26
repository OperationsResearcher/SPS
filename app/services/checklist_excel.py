# -*- coding: utf-8 -*-
"""Kontrol listesi Excel içe/dışa aktarma.

İçe aktarma ProjeKontrol.xlsx düzenini bekler:
    SÜREÇ (birleşik hücre, aşama) | [alt grup: Uygulama/Çıktı] | ADIM | SORUMLU | DURUM
SORUMLU ve DURUM okunmaz: sorumlular proje içinde maddeye atanır (kişi / süreç / proje ekibi).
Sütunlar başlık satırından bulunur; aşama ve grup birleşik hücrelerde yalnızca ilk satırda
yazılı olduğu için aşağı taşınır.
"""
from __future__ import annotations

import io
import re
from typing import Iterable

_TR = str.maketrans("ÇĞİIÖŞÜçğıiöşü", "CGIIOSUcgiiosu")


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").translate(_TR)).strip().upper()


def _clean(s) -> str:
    """Boşlukları toparlar, baştaki madde işaretini (-, –, •) atar."""
    t = re.sub(r"\s+", " ", str(s or "")).strip()
    return re.sub(r"^[-–•*]+\s*", "", t).strip()


_PERIOD_RE = re.compile(r"(\d+)\s*haftada\s*bir", re.I)
_PERIOD_DAY_RE = re.compile(r"(\d+)\s*g[uü]nde\s*bir", re.I)


def detect_period_days(text: str) -> int | None:
    """'2 haftada bir', '10 günde bir' gibi ifadelerden periyot (gün) çıkarır."""
    m = _PERIOD_RE.search(text or "")
    if m:
        return int(m.group(1)) * 7
    m = _PERIOD_DAY_RE.search(text or "")
    if m:
        return int(m.group(1))
    return None


def parse_workbook(data: bytes) -> list[dict]:
    """xlsx baytlarından madde listesi: [{stage, item_group, text, period_days}].

    Hata: ValueError (başlık satırı bulunamazsa vb.) — mesaj kullanıcıya gösterilir.
    """
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # bozuk / xlsx olmayan dosya
        raise ValueError("Dosya okunamadı; geçerli bir .xlsx dosyası yükleyin.") from exc
    ws = wb.worksheets[0]

    header_row = col_text = col_stage = None
    for r in range(1, min(ws.max_row, 15) + 1):
        for c in range(1, ws.max_column + 1):
            if _norm(ws.cell(r, c).value) == "ADIM":
                header_row, col_text = r, c
                break
        if header_row:
            break
    if not header_row:
        raise ValueError("'ADIM' başlıklı sütun bulunamadı; dosya ProjeKontrol düzeninde olmalı.")
    for c in range(1, ws.max_column + 1):
        h = _norm(ws.cell(header_row, c).value)
        if h in ("SUREC", "ASAMA"):
            col_stage = c
    col_stage = col_stage or 1
    col_group = col_stage + 1 if col_stage + 1 < col_text else None

    items: list[dict] = []
    stage = group = None
    for r in range(header_row + 1, ws.max_row + 1):
        st = _clean(ws.cell(r, col_stage).value)
        if st:
            stage, group = st, None
        if col_group:
            g = _clean(ws.cell(r, col_group).value)
            if g:
                group = g
        text = _clean(ws.cell(r, col_text).value)
        if not text:
            continue
        if not stage:
            raise ValueError(f"Satır {r}: madde için aşama (SÜREÇ) bulunamadı.")
        items.append({
            "stage": stage,
            "item_group": group,
            "text": text,
            "period_days": detect_period_days(text),
        })
    if not items:
        raise ValueError("Dosyada içe aktarılacak madde bulunamadı.")
    return items


def template_rows(items: Iterable) -> list[dict]:
    return [
        {"stage": i.stage, "item_group": i.item_group, "text": i.text,
         "period_days": i.period_days}
        for i in items
    ]
