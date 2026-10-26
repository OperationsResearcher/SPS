"""Proje kontrol listesi: şablon kopyası, çift onay, kilitli kapanış, rol ataması."""
from datetime import date

import pytest
from werkzeug.security import generate_password_hash

from app.models import db
from app.models.core import Role, Tenant, User
from app.models.portfolio_project import Project, project_leaders
from app.models.project_checklist import (
    ChecklistTemplate,
    ChecklistTemplateItem,
    ProjectChecklistItem,
    STATUS_APPROVED,
    STATUS_DONE,
    STATUS_PENDING,
)
from app.services import project_checklist_service as svc


@pytest.fixture
def ck(app):
    t = Tenant(name="T", short_name="t", is_active=True)
    r_admin = Role(name="tenant_admin", description="ta")
    r_std = Role(name="standard_user", description="su")
    db.session.add_all([t, r_admin, r_std])
    db.session.flush()

    def mk(email, role):
        u = User(email=email, first_name=email.split("@")[0], last_name="x", tenant_id=t.id,
                 role_id=role.id, password_hash=generate_password_hash("x"), is_active=True)
        db.session.add(u)
        return u

    lider, lider2, uzman, yabanci = (mk("lider@t.local", r_std), mk("lider2@t.local", r_std),
                                     mk("uzman@t.local", r_std), mk("yabanci@t.local", r_std))
    db.session.flush()
    tpl = ChecklistTemplate(tenant_id=t.id, year=2026, name="Danışmanlık")
    db.session.add(tpl)
    db.session.flush()
    db.session.add_all([
        ChecklistTemplateItem(template_id=tpl.id, stage="Teşhis", item_group="Uygulama", text="A",
                              sort_order=1),
        ChecklistTemplateItem(template_id=tpl.id, stage="Teşhis", item_group="Çıktı", text="B",
                              sort_order=2, evidence_required=True),
        ChecklistTemplateItem(template_id=tpl.id, stage="Tedavi", text="C",
                              sort_order=3, period_days=14),
        ChecklistTemplateItem(template_id=tpl.id, stage="Tedavi", text="Pasif", sort_order=4, is_active=False),
    ])
    proj = Project(name="P", kurum_id=t.id, manager_id=lider.id, start_date=date(2026, 3, 1))
    db.session.add(proj)
    db.session.flush()
    db.session.execute(project_leaders.insert().values(project_id=proj.id, user_id=lider.id))
    db.session.execute(project_leaders.insert().values(project_id=proj.id, user_id=lider2.id))
    db.session.commit()
    return {"t": t, "p": proj, "lider": lider, "lider2": lider2, "uzman": uzman, "yabanci": yabanci}


def test_sablon_projeye_kopyalanir_pasif_haric(ck):
    assert svc.copy_template_to_project(ck["p"]) == 3
    assert svc.copy_template_to_project(ck["p"]) == 0  # ikinci kez dokunmaz
    items = svc.active_items(ck["p"].id)
    assert [i.text for i in items] == ["A", "B", "C"]
    assert items[2].next_reminder_at is not None  # periyodik madde


def test_sablon_yoksa_kopya_yok(ck):
    ck["p"].start_date = date(2024, 1, 1)  # 2024 şablonu yok
    db.session.commit()
    assert svc.copy_template_to_project(ck["p"]) == 0


def test_cift_onay_akisi(ck):
    svc.copy_template_to_project(ck["p"])
    p, lider, lider2, uzman, yab = ck["p"], ck["lider"], ck["lider2"], ck["uzman"], ck["yabanci"]
    a = svc.active_items(p.id)[0]
    with pytest.raises(svc.ChecklistError) as e:  # atanmamış yabancı işaretleyemez
        svc.mark_done(p, a, yab)
    assert e.value.status == 403
    svc.set_assignments(p, lider, [a.id], [{"type": "user", "id": uzman.id}])
    db.session.refresh(a)
    svc.mark_done(p, a, uzman)
    assert a.status == STATUS_DONE and a.done_by == uzman.id
    with pytest.raises(svc.ChecklistError):  # uzman onaylayamaz (lider değil)
        svc.approve(p, a, uzman)
    svc.approve(p, a, lider)
    assert a.status == STATUS_APPROVED and a.approved_by == lider.id


def test_yapan_kendini_onaylayamaz(ck):
    svc.copy_template_to_project(ck["p"])
    p, lider, lider2 = ck["p"], ck["lider"], ck["lider2"]
    a = svc.active_items(p.id)[0]
    svc.mark_done(p, a, lider)  # lider kendisi işaretledi
    with pytest.raises(svc.ChecklistError) as e:
        svc.approve(p, a, lider)
    assert e.value.status == 403
    svc.approve(p, a, lider2)  # ikinci yetkili onaylar
    assert a.status == STATUS_APPROVED


def test_kanit_zorunlu_madde(ck):
    svc.copy_template_to_project(ck["p"])
    p, lider = ck["p"], ck["lider"]
    b = svc.active_items(p.id)[1]
    with pytest.raises(svc.ChecklistError) as e:
        svc.mark_done(p, b, lider)
    assert e.value.status == 422
    from app.models.project_checklist import ProjectChecklistEvidence
    db.session.add(ProjectChecklistEvidence(item_id=b.id, uploader_id=lider.id, kind="link",
                                            title="k", url="https://x.test"))
    db.session.flush()
    db.session.refresh(b)
    svc.mark_done(p, b, lider)
    assert b.status == STATUS_DONE


def test_kilitli_kapanis(ck):
    svc.copy_template_to_project(ck["p"])
    p, lider, lider2 = ck["p"], ck["lider"], ck["lider2"]
    with pytest.raises(svc.ChecklistError) as e:
        svc.complete_project(p, lider)
    assert e.value.status == 422
    from app.models.project_checklist import ProjectChecklistEvidence
    for it in svc.active_items(p.id):
        if it.evidence_required:
            db.session.add(ProjectChecklistEvidence(item_id=it.id, uploader_id=lider.id, kind="link",
                                                    title="k", url="https://x.test"))
            db.session.flush()
            db.session.refresh(it)
        svc.mark_done(p, it, lider)
        svc.approve(p, it, lider2)
    assert svc.progress(p.id)["pct"] == 100
    svc.complete_project(p, lider)
    assert p.completed_at is not None
    with pytest.raises(svc.ChecklistError) as e:  # tamamlanınca kilitli
        svc.add_item(p, lider, stage="X", text="Y")
    assert e.value.status == 423
    svc.reopen_project(p, lider)
    svc.add_item(p, lider, stage="X", text="Y")
    assert svc.completion_state(p.id)["can_complete"] is False


def test_madde_ekle_sil_ilerleme(ck):
    svc.copy_template_to_project(ck["p"])
    p, lider, yab = ck["p"], ck["lider"], ck["yabanci"]
    with pytest.raises(svc.ChecklistError):
        svc.add_item(p, yab, stage="X", text="Y")
    n = svc.add_item(p, lider, stage="Yeni", text="Ek madde")
    items = svc.active_items(p.id)
    assert n in items and len(items) == 4
    svc.remove_item(p, n, lider)
    assert len(svc.active_items(p.id)) == 3
    pr = svc.progress(p.id)
    assert pr["total"] == 3 and pr["approved"] == 0 and [s["stage"] for s in pr["stages"]] == ["Teşhis", "Tedavi"]
    assert svc.progress_map([p.id])[p.id]["total"] == 3


def test_yeni_proje_olusturma_sablonu_kopyalar(app, client, ck):
    """project_new route'u şablonu otomatik kopyalar."""
    admin = User.query.filter_by(email="lider@t.local").first()
    r = Role.query.filter_by(name="tenant_admin").first()
    admin.role_id = r.id
    db.session.commit()
    with client.session_transaction() as s:
        s["_user_id"] = str(admin.id)
        s["_fresh"] = True
        s["sv"] = 0
    rv = client.post("/k-plan/project/new", data={
        "name": "Yeni", "start_date": "2026-05-01", "leaders": [str(admin.id)],
    })
    assert rv.status_code in (302, 303), rv.data[:300]
    new_id = int(rv.headers["Location"].rstrip("/").rsplit("/", 1)[1])
    assert new_id != ck["p"].id
    assert ProjectChecklistItem.query.filter_by(project_id=new_id).count() == 3


# ── Aşama 2: şablon yönetimi ──────────────────────────────────────────────────
def _login(client, user):
    with client.session_transaction() as s:
        s["_user_id"] = str(user.id)
        s["_fresh"] = True
        s["sv"] = 0


def _xlsx_bytes():
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Başlık"])
    ws.append(["SÜREÇ", None, "ADIM", "SORUMLU", "DURUM"])
    ws.append(["GENEL", None, "- 2 haftada bir toplantı yapılacak", "Uzman", False])
    ws.append([None, None, "-Klasör düzeni", "Pazarlama", False])
    ws.append(["TEŞHİS", "Uygulama", "Süreç akışı çizilecek", "Uzman", False])
    ws.append([None, "Çıktı", "Rapor hazırlanacak", "Uzman", False])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_ayristirici():
    from app.services.checklist_excel import parse_workbook
    rows = parse_workbook(_xlsx_bytes())
    assert [(r["stage"], r["item_group"], r["text"]) for r in rows] == [
        ("GENEL", None, "2 haftada bir toplantı yapılacak"),
        ("GENEL", None, "Klasör düzeni"),
        ("TEŞHİS", "Uygulama", "Süreç akışı çizilecek"),
        ("TEŞHİS", "Çıktı", "Rapor hazırlanacak"),
    ]
    assert rows[0]["period_days"] == 14 and rows[1]["period_days"] is None
    with pytest.raises(ValueError):
        parse_workbook(b"bozuk")


def test_sablon_api_yetki_ve_crud(client, ck):
    r_admin = Role.query.filter_by(name="tenant_admin").first()
    yonetici = ck["lider2"]
    yonetici.role_id = r_admin.id
    db.session.commit()

    _login(client, ck["uzman"])  # standart kullanıcı
    assert client.get("/k-plan/project/api/checklist-template?year=2027").status_code == 403
    assert client.get("/k-plan/project/checklist-template").status_code == 302

    _login(client, yonetici)
    assert client.get("/k-plan/project/checklist-template").status_code == 200
    j = client.get("/k-plan/project/api/checklist-template?year=2026").get_json()
    assert [i["text"] for i in j["items"]] == ["A", "B", "C"] and 2026 in j["years"]

    j = client.post("/k-plan/project/api/checklist-template/item",
                    json={"year": 2027, "stage": "S1", "text": "Yeni",
                          "period_days": 7, "evidence_required": True}).get_json()
    assert j["success"] and len(j["items"]) == 1 and j["items"][0]["period_days"] == 7
    iid = j["items"][0]["id"]
    j = client.post("/k-plan/project/api/checklist-template/item",
                    json={"year": 2027, "id": iid, "stage": "S1", "text": "Güncel"}).get_json()
    assert j["items"][0]["text"] == "Güncel" and j["items"][0]["period_days"] is None
    assert client.post("/k-plan/project/api/checklist-template/item",
                       json={"year": 2027, "stage": "", "text": "x"}).status_code == 400
    assert client.post("/k-plan/project/api/checklist-template/item",
                       json={"year": 2027, "stage": "S", "text": "x", "period_days": 999}).status_code == 400

    # yıl kopyası (2026 → 2028) ve dolu hedefe kopya reddi
    j = client.post("/k-plan/project/api/checklist-template/copy", json={"from_year": 2026, "to_year": 2028}).get_json()
    assert j["copied"] == 3
    assert client.post("/k-plan/project/api/checklist-template/copy",
                       json={"from_year": 2026, "to_year": 2028}).status_code == 409

    j = client.post(f"/k-plan/project/api/checklist-template/item/{iid}/delete", json={"year": 2027}).get_json()
    assert j["items"] == []


def test_sablon_excel_ice_aktar_ve_yeniden_sirala(client, ck):
    import io
    yonetici = ck["lider2"]
    yonetici.role_id = Role.query.filter_by(name="tenant_admin").first().id
    db.session.commit()
    _login(client, yonetici)
    rv = client.post("/k-plan/project/api/checklist-template/import",
                     data={"file": (io.BytesIO(_xlsx_bytes()), "k.xlsx"), "year": "2030", "mode": "replace"},
                     content_type="multipart/form-data")
    j = rv.get_json()
    assert rv.status_code == 200 and j["imported"] == 4 and len(j["items"]) == 4
    assert client.post("/k-plan/project/api/checklist-template/import",
                       data={"file": (io.BytesIO(b"x"), "k.txt"), "year": "2030"},
                       content_type="multipart/form-data").status_code == 400
    ids = [i["id"] for i in j["items"]]
    j = client.post("/k-plan/project/api/checklist-template/reorder",
                    json={"year": 2030, "ids": list(reversed(ids))}).get_json()
    assert [i["id"] for i in j["items"]] == list(reversed(ids))
    # sonuna ekle modu
    j = client.post("/k-plan/project/api/checklist-template/import",
                    data={"file": (io.BytesIO(_xlsx_bytes()), "k.xlsx"), "year": "2030", "mode": "append"},
                    content_type="multipart/form-data").get_json()
    assert len(j["items"]) == 8


# ── Aşama 3: proje içi API ───────────────────────────────────────────────────
def _api(ck):
    return f"/k-plan/project/{ck['p'].id}/api/checklist"


def test_proje_checklist_sayfasi_ve_api_akisi(client, ck):
    import io
    svc.copy_template_to_project(ck["p"])
    db.session.commit()
    from app.models.portfolio_project import project_members
    db.session.execute(project_members.insert().values(project_id=ck["p"].id, user_id=ck["uzman"].id))
    db.session.commit()
    base = _api(ck)

    _login(client, ck["yabanci"])  # projede yok
    assert client.get(base).status_code == 403
    assert client.get(f"/k-plan/project/{ck['p'].id}/views/checklist").status_code == 302

    _login(client, ck["lider"])
    _r = client.get(f"/k-plan/project/{ck['p'].id}/views/checklist")
    _k = client.get(f"/k-plan/project/{ck['p'].id}/views/kanban")
    assert _r.status_code == _k.status_code, (_r.status_code, _k.status_code)
    j = client.get(base).get_json()
    assert j["can_manage"] and j["progress"]["total"] == 3 and j["project"]["team_size"] >= 2
    assert len(j["users"]) >= 4 and j["processes"] == []
    assert j["completion"]["can_complete"] is False
    ids = {i["text"]: i["id"] for i in j["items"]}

    # kişi ataması → uzman kendi maddesini işaretler
    j = client.post(f"{base}/assign", json={"item_ids": [ids["A"], ids["B"]], "targets": [{"type": "user", "id": ck["uzman"].id}]}).get_json()
    assert j["success"]
    _login(client, ck["uzman"])
    j = client.get(base).get_json()
    a = next(i for i in j["items"] if i["text"] == "A")
    assert a["can_done"] and not a["can_approve"] and j["can_manage"] is False
    assert client.post(f"{base}/assign", json={"item_ids": [ids["A"]], "targets": []}).status_code == 403
    assert client.post(f"{base}/item/{ids['A']}/done").get_json()["success"]
    assert client.post(f"{base}/item/{ids['A']}/approve").status_code == 403  # uzman onaylayamaz

    # kanıt zorunlu madde: kanıtsız 422, bağlantı + dosya ile geçer
    assert client.post(f"{base}/item/{ids['B']}/done").status_code == 422
    r = client.post(f"{base}/item/{ids['B']}/evidence", json={"url": "javascript:alert(1)"})
    assert r.status_code == 400
    r = client.post(f"{base}/item/{ids['B']}/evidence", json={"url": "https://x.test/k", "title": "Foto"})
    assert r.status_code == 200
    r = client.post(f"{base}/item/{ids['B']}/evidence",
                    data={"file": (io.BytesIO(b"%PDF-1.4 test"), "kanit.pdf")}, content_type="multipart/form-data")
    assert r.status_code == 200
    bad = client.post(f"{base}/item/{ids['B']}/evidence",
                      data={"file": (io.BytesIO(b"MZ"), "x.exe")}, content_type="multipart/form-data")
    assert bad.status_code == 400
    j = client.post(f"{base}/item/{ids['B']}/done").get_json()
    b = next(i for i in j["items"] if i["text"] == "B")
    assert b["status"] == "Yapıldı" and len(b["evidences"]) == 2
    dl = next(e for e in b["evidences"] if e["download_url"])["download_url"]
    assert client.get(dl).status_code == 200
    _login(client, ck["yabanci"])
    assert client.get(dl).status_code in (403, 404)

    # lider2 onaylar, tamamlama kilidi
    _login(client, ck["lider2"])
    assert client.post(f"{base}/complete").status_code == 422
    for k in ("A", "B"):
        assert client.post(f"{base}/item/{ids[k]}/approve").get_json()["success"]
    j = client.post(f"{base}/item/{ids['C']}/done").get_json()
    assert client.post(f"{base}/item/{ids['C']}/approve").status_code == 403  # yapan = onaylayan olamaz
    _login(client, ck["lider"])
    assert client.post(f"{base}/item/{ids['C']}/approve").get_json()["progress"]["pct"] == 100
    j = client.post(f"{base}/complete").get_json()
    assert j["project"]["completed_at"]
    assert client.post(f"{base}/item", json={"stage": "X", "text": "Y"}).status_code == 423
    assert client.post(f"{base}/reopen-project").get_json()["project"]["completed_at"] is None


def test_proje_checklist_sayfasi_render(client, ck):
    ck["lider2"].role_id = Role.query.filter_by(name="tenant_admin").first().id
    db.session.commit()
    _login(client, ck["lider2"])
    rv = client.get(f"/k-plan/project/{ck['p'].id}/views/checklist")
    assert rv.status_code == 200
    html = rv.data.decode("utf-8")
    assert "ckp-root" in html and "project_checklist.js" in html and "<style" not in html.split("</head>")[1]
    assert "Kontrol Listesi" in client.get(f"/k-plan/project/{ck['p'].id}/views/kanban").data.decode("utf-8")  # sekme menüsü


def test_liste_ve_detayda_ilerleme(client, ck):
    ck["lider2"].role_id = Role.query.filter_by(name="tenant_admin").first().id
    svc.copy_template_to_project(ck["p"])
    items = svc.active_items(ck["p"].id)
    svc.mark_done(ck["p"], items[0], ck["lider"])
    svc.approve(ck["p"], items[0], ck["lider2"])
    db.session.commit()
    _login(client, ck["lider2"])
    html = client.get("/k-plan/project").data.decode("utf-8")
    assert "ck-progress-native" in html and "%33" in html and "1/3" in html
    html = client.get(f"/k-plan/project/{ck['p'].id}").data.decode("utf-8")
    assert "project_detail.kontrol_listesi_ozeti" in html and "%33" in html
    assert svc.progress_map([ck["p"].id])[ck["p"].id]["pct"] == 33


def test_periyodik_hatirlatma(ck):
    from datetime import datetime, timedelta
    from app.models.core import Notification
    p = ck["p"]
    svc.copy_template_to_project(p)
    c = next(i for i in svc.active_items(p.id) if i.text == "C")  # 14 günlük periyot
    svc.set_assignments(p, ck["lider"], [c.id], [{"type": "user", "id": ck["uzman"].id}])
    db.session.commit()
    db.session.refresh(c)
    now = datetime(2026, 3, 10, 10, 0)  # proje başlangıcı 1 Mart → ilk hatırlatma 15 Mart
    assert svc.run_due_reminders(now) == 0  # sonraki hatırlatma henüz gelmedi
    c.next_reminder_at = now - timedelta(hours=1)
    db.session.commit()
    assert svc.run_due_reminders(now) == 1
    db.session.commit()
    notif = Notification.query.filter_by(user_id=ck["uzman"].id, notification_type="checklist_reminder").all()
    assert len(notif) == 1 and "C" in notif[0].message
    assert c.next_reminder_at == now + timedelta(days=14) and c.last_reminded_at == now
    assert svc.run_due_reminders(now) == 0  # aynı periyotta tekrar yok
    # onaylı madde ve tamamlanmış proje hatırlatmaz
    c.next_reminder_at = now - timedelta(hours=1)
    c.status = STATUS_APPROVED
    db.session.commit()
    assert svc.run_due_reminders(now) == 0
    c.status = STATUS_PENDING
    p.completed_at = now
    db.session.commit()
    assert svc.run_due_reminders(now) == 0
    p.completed_at = None
    p.is_archived = True
    db.session.commit()
    assert svc.run_due_reminders(now) == 0


def test_excel_ve_pdf_cikti(client, ck):
    import io
    import openpyxl
    p = ck["p"]
    svc.copy_template_to_project(p)
    items = svc.active_items(p.id)
    svc.mark_done(p, items[0], ck["lider"])
    svc.approve(p, items[0], ck["lider2"])
    db.session.commit()
    _login(client, ck["lider"])
    rv = client.get(f"/k-plan/project/{p.id}/api/checklist/export.xlsx")
    assert rv.status_code == 200 and "spreadsheetml" in rv.mimetype
    assert "attachment" in rv.headers["Content-Disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(rv.data)).active
    texts = [ws.cell(r, 3).value for r in range(4, 4 + 3)]
    assert texts == ["A", "B", "C"]
    assert "☑" in ws.cell(4, 5).value and "☐" in ws.cell(5, 5).value
    assert "%33" in ws["A2"].value
    # Aşama "Teşhis" (A,B) birleşik ve gruplar ayrı; "Tedavi" tek satır
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert "A4:A5" in merged
    # ayrıştırıcıyla geri okunabilir (orijinal düzen)
    from app.services.checklist_excel import parse_workbook
    back = parse_workbook(rv.data)
    assert [b["text"] for b in back] == ["A", "B", "C"] and back[0]["stage"] == "Teşhis"
    rv = client.get(f"/k-plan/project/{p.id}/api/checklist/export.pdf")
    assert rv.status_code == 200 and rv.data[:4] == b"%PDF" and len(rv.data) > 1500
    assert client.get(f"/k-plan/project/{p.id}/api/checklist/export.csv").status_code == 404
    _login(client, ck["yabanci"])
    assert client.get(f"/k-plan/project/{p.id}/api/checklist/export.pdf").status_code == 403


def test_atama_hedefleri_surec_proje_kullanici(client, ck):
    """Maddeye süreç (tüm üyeleri), proje ekibi ve tek tek kullanıcılar atanabilir."""
    from app.models.process import Process, process_members
    from app.models.portfolio_project import project_members
    p, lider, lider2, uzman, yab = ck["p"], ck["lider"], ck["lider2"], ck["uzman"], ck["yabanci"]
    svc.copy_template_to_project(p)
    ek = User(email="ekip@t.local", first_name="ekip", last_name="x", tenant_id=ck["t"].id,
              role_id=lider.role_id, password_hash="x", is_active=True)
    db.session.add(ek)
    db.session.flush()
    db.session.execute(project_members.insert().values(project_id=p.id, user_id=ek.id))
    proc = Process(name="Satis", code="SR1", tenant_id=ck["t"].id, is_active=True)
    db.session.add(proc)
    db.session.flush()
    db.session.execute(process_members.insert().values(process_id=proc.id, user_id=yab.id))
    db.session.commit()
    a, b, c = svc.active_items(p.id)

    # süreç → sürecin üyesi (yab) maddeyi işaretleyebilir
    assert not svc.can_work(yab, p, a)
    svc.set_assignments(p, lider, [a.id], [{"type": "process", "id": proc.id}])
    assert svc.user_is_assigned(yab, p, a) and svc.expand_assignees(p, a) == {yab.id}
    svc.mark_done(p, a, yab)
    assert a.status == STATUS_DONE

    # proje → proje ekibi (lider, lider2, ekip üyesi); yab ekipte değil
    svc.set_assignments(p, lider, [b.id], [{"type": "project"}])
    team = svc.expand_assignees(p, b)
    assert {lider.id, lider2.id, ek.id} <= team and yab.id not in team and svc.can_work(ek, p, b)

    # birden çok hedef + çoklu madde; 'ekle' modu mevcutları korur
    svc.set_assignments(p, lider, [c.id], [{"type": "user", "id": uzman.id}])
    svc.set_assignments(p, lider, [c.id, b.id], [{"type": "process", "id": proc.id}], replace=False)
    labels = {x["type"] for x in svc.assignment_labels(p, c)}
    assert labels == {"user", "process"} and svc.expand_assignees(p, c) == {uzman.id, yab.id}
    # 'değiştir' modu eskileri siler
    svc.set_assignments(p, lider, [c.id], [{"type": "user", "id": ek.id}])
    assert svc.expand_assignees(p, c) == {ek.id}

    # doğrulamalar: yetki, yabancı kurum süreci/kullanıcısı, geçersiz hedef
    with pytest.raises(svc.ChecklistError) as e:
        svc.set_assignments(p, uzman, [a.id], [{"type": "user", "id": uzman.id}])
    assert e.value.status == 403
    other = Tenant(name="Diger", short_name="d", is_active=True)
    db.session.add(other)
    db.session.flush()
    foreign_user = User(email="f@d.local", first_name="f", last_name="f", tenant_id=other.id,
                        role_id=lider.role_id, password_hash="x", is_active=True)
    foreign_proc = Process(name="Yabanci", tenant_id=other.id, is_active=True)
    db.session.add_all([foreign_user, foreign_proc])
    db.session.flush()
    for bad in ({"type": "user", "id": foreign_user.id}, {"type": "process", "id": foreign_proc.id},
                {"type": "team", "id": 1}, {"type": "user", "id": "x"}):
        with pytest.raises(svc.ChecklistError) as e:
            svc.set_assignments(p, lider, [a.id], [bad])
        assert e.value.status == 400
    with pytest.raises(svc.ChecklistError):
        svc.set_assignments(p, lider, [], [{"type": "project"}])


def test_atama_api_ve_bana_atananlar(client, ck):
    from app.models.process import Process, process_members
    p = ck["p"]
    svc.copy_template_to_project(p)
    proc = Process(name="Uretim", code="UR1", tenant_id=ck["t"].id, is_active=True)
    db.session.add(proc)
    db.session.flush()
    db.session.execute(process_members.insert().values(process_id=proc.id, user_id=ck["yabanci"].id))
    db.session.commit()
    base = _api(ck)
    _login(client, ck["lider"])
    j = client.get(base).get_json()
    assert [x["id"] for x in j["processes"]] == [proc.id] and j["processes"][0]["name"].startswith("[UR1]")
    ids = [i["id"] for i in j["items"]]
    r = client.post(f"{base}/assign", json={"item_ids": ids[:2], "targets": [{"type": "process", "id": proc.id},
                                                                          {"type": "project"}]})
    assert r.status_code == 200
    items = {i["id"]: i for i in r.get_json()["items"]}
    assert {a["type"] for a in items[ids[0]]["assignments"]} == {"process", "project"}
    assert items[ids[2]]["assignments"] == []
    assert client.post(f"{base}/assign", json={"item_ids": ids, "targets": [{"type": "process", "id": 99999}]}).status_code == 400
    assert client.post(f"{base}/assign", json={"item_ids": "x"}).status_code == 400
    # sürecin üyesi 'Bana atananlar' görür; projede olmasa da (süreç üzerinden) erişimi proje izinlerine bağlı
    from app.models.portfolio_project import project_members
    db.session.execute(project_members.insert().values(project_id=p.id, user_id=ck["yabanci"].id))
    db.session.commit()
    _login(client, ck["yabanci"])
    j = client.get(base).get_json()
    assert [i["mine"] for i in j["items"]] == [True, True, False]
    assert j["users"] == [] and j["processes"] == []  # yönetici olmayana seçenek listesi verilmez
